import os
import openpyxl


def load_piezometer_readings(file_path, sheet_name="P01DS1"):
    """
    Load piezometer pressure readings from an ACCRD Excel sheet.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    workbook = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")

    sheet = workbook[sheet_name]

    readings = []

    # Actual observation data starts from row 16
    for row in sheet.iter_rows(min_row=16, values_only=True):

        observation_date = row[0]
        observation_time = row[1]
        pressure = row[4]

        if observation_date is None or pressure is None:
            continue

        try:
            pressure = float(pressure)
        except (TypeError, ValueError):
            continue

        readings.append({
            "instrument_id": sheet_name,
            "metric": "pore_pressure",
            "value": pressure,
            "date": observation_date,
            "time": observation_time,
            "unit": "MPa"
        })

    return readings
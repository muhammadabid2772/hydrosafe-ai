from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl


def load_piezometer_readings(file_path: str | Path, sheet_name: str = "P01DS1") -> list[dict]:
    """Load the ACCRD piezometer layout used in Member 3's handoff."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        sheet = workbook[sheet_name]
        readings = []
        for row in sheet.iter_rows(min_row=16, values_only=True):
            observation_date, observation_time, pressure = row[0], row[1], row[4]
            if observation_date is None or pressure is None:
                continue
            try:
                pressure = float(pressure)
            except (TypeError, ValueError):
                continue
            timestamp = observation_date
            if isinstance(observation_date, datetime) and observation_time is not None:
                timestamp = datetime.combine(observation_date.date(), observation_time)
            readings.append({
                "instrument_id": sheet_name,
                "metric": "pore_pressure",
                "value": pressure,
                "timestamp": timestamp.isoformat() if isinstance(timestamp, datetime) else str(timestamp),
                "unit": "MPa",
            })
        return readings
    finally:
        workbook.close()

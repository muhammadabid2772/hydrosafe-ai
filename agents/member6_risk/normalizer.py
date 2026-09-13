from __future__ import annotations

import json
import math
import re
import warnings
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


STRUCTURE_BY_FOLDER = {"ACCRD": "ACCRD", "PH": "POWER_HOUSE", "PI": "POWER_INTAKE"}
DATE_WORDS = ("date", "日期", "观测时间", "survey time")
TIME_WORDS = ("time", "时间")
UNIT_RE = re.compile(r"(?:\(|（)([^()（）]{1,12})(?:\)|）)")
INSTRUMENT_RE = re.compile(
    r"^(?:BV|E|ID|J|P|S|TC|WE|GB|M|R|D|TP|DS|CZ)[- ]?\d+[A-Z0-9-]*(?:\(.*\))?$",
    re.IGNORECASE,
)
SUMMARY_WORDS = (
    "安全鉴定", "report", "chart", "graph", "过程线", "成果统计", "年报", "周报", "月报",
    "water level", "水位", "填筑高程", "plan", "sheet", "蓄水快报", "data input",
)
METRIC_PRIORITIES = {
    "E": ("applied pressure", "土压力"),
    "ID": ("累计位移量", "absolute displacement", "绝对位移"),
    "TC": ("坝体累计沉降量", "absolute settlement", "绝对沉降", "累计沉降"),
    "CZ": ("△s", "δs", "累计位移", "累积位移"),
    "P": ("pressure", "渗透压力", "water el", "水头高程", "水位高程"),
    "BV": ("water level elevation", "水位高程", "water level", "水位"),
    "R": ("stress", "应力", "force", "拉力"),
    "D": ("force", "荷载", "拉力", "load"),
    "M": ("calculated", "计算值", "settlement", "沉降量", "displacement", "位移", "deformation", "变形"),
    "WE": ("q(l/s)", "seepage flow", "渗流量", "q(m3/h)"),
    "TP": ("displacement", "位移量", "offset", "偏移量"),
    "BM": ("settlement", "沉降量", "elevation", "高程"),
    "J": ("joint", "opening", "开度"),
    "S": ("strain", "应变"),
}


@dataclass
class NormalizedReading:
    structure: str
    instrument_type: str
    instrument_id: str
    timestamp: str
    metric: str
    value: float
    unit: str | None
    source_file: str
    source_sheet: str


def _text(value: object) -> str:
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _instrument_type(instrument_id: str) -> str:
    token = re.match(r"[A-Za-z]+", instrument_id.replace("-", ""))
    return token.group(0).upper() if token else "UNKNOWN"


def _structure(path: Path, sheet: str) -> str | None:
    parts = {part.upper() for part in path.parts}
    for folder, value in STRUCTURE_BY_FOLDER.items():
        if folder in parts:
            return value
    upper = sheet.upper()
    if upper.endswith("AD") or "ACCRD" in path.name.upper():
        return "ACCRD"
    if "GPH" in upper or upper.endswith("PH") or "GPH" in path.name.upper():
        return "POWER_HOUSE"
    if upper.endswith("PI") or upper.endswith("PIT") or "PIT" in path.name.upper():
        return "POWER_INTAKE"
    return None


def _sheet_instrument_id(sheet: str) -> str | None:
    cleaned = re.sub(r"（.*?）", "", sheet).strip()
    return cleaned if INSTRUMENT_RE.match(cleaned) else None


def _find_header(rows: list[list[object]]) -> tuple[int, int] | None:
    preferred = ("observation date", "观测日期", "survey time")
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            value_text = _text(value).lower()
            if any(word in value_text for word in preferred):
                return row_index, col_index
    for row_index, row in enumerate(rows):
        text = [_text(value).lower() for value in row]
        date_columns = [index for index, value in enumerate(text) if any(word in value for word in DATE_WORDS)]
        has_time = any(any(word in value for word in TIME_WORDS) for value in text)
        if date_columns and has_time and not any(word in value for value in text for word in ("installation", "install date", "埋设")):
            return row_index, date_columns[0]
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            value_text = _text(value).lower()
            if any(word in value_text for word in DATE_WORDS):
                return row_index, col_index
    return None


def _combine_headers(rows: list[list[object]], header_row: int, column: int) -> str:
    values = []
    for index in range(header_row, min(len(rows), header_row + 2)):
        value = _text(rows[index][column] if column < len(rows[index]) else None)
        if value and value not in values:
            values.append(value)
    return " / ".join(values)


def _timestamp(value: object, time_value: object = None) -> datetime | None:
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, time())
    elif isinstance(value, str):
        result = None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%Y.%m.%d"):
            try:
                result = datetime.strptime(value.strip(), fmt)
                break
            except ValueError:
                pass
        if result is None:
            return None
    elif isinstance(value, (int, float)) and 20_000 <= value <= 60_000:
        result = from_excel(value)
    else:
        return None
    if isinstance(time_value, time):
        result = datetime.combine(result.date(), time_value)
    return result


def _numeric(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def _primary_column(headers: dict[int, str], instrument_type: str) -> int | None:
    for word in METRIC_PRIORITIES.get(instrument_type, ()):
        for column, header in headers.items():
            if word in header.lower():
                return column
    ignored = ("digit", "模数", "temperature", "温度", "time", "日期", "remark", "备注")
    candidates = [column for column, header in headers.items() if header and not any(x in header.lower() for x in ignored)]
    return candidates[-1] if candidates else None


def iter_workbook_readings(path: str | Path, root: str | Path | None = None) -> Iterator[NormalizedReading]:
    path = Path(path)
    root_path = Path(root) if root else None
    relative = str(path.relative_to(root_path)) if root_path and path.is_relative_to(root_path) else path.name
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for worksheet in workbook.worksheets:
            instrument_id = _sheet_instrument_id(worksheet.title)
            if not instrument_id or any(word in worksheet.title.lower() for word in SUMMARY_WORDS):
                continue
            structure = _structure(path, worksheet.title)
            if structure is None:
                continue
            instrument_type = _instrument_type(instrument_id)
            max_column = min(worksheet.max_column, 40)
            rows = [list(row) for row in worksheet.iter_rows(max_col=max_column, values_only=True)]
            found = _find_header(rows[:30])
            if not found:
                continue
            header_row, date_column = found
            headers = {column: _combine_headers(rows, header_row, column) for column in range(max_column)}
            time_column = next(
                (column for column, header in headers.items() if any(word in header.lower() for word in TIME_WORDS)), None
            )
            primary_column = _primary_column(headers, instrument_type)
            if primary_column is None:
                continue
            metric = headers[primary_column] or f"column_{primary_column + 1}"
            unit_match = UNIT_RE.search(metric)
            unit = unit_match.group(1) if unit_match else None
            for row in rows[header_row + 1 :]:
                if date_column >= len(row):
                    continue
                observed_at = _timestamp(row[date_column], row[time_column] if time_column is not None else None)
                value = _numeric(row[primary_column] if primary_column < len(row) else None)
                if observed_at is None or value is None:
                    continue
                yield NormalizedReading(
                    structure=structure, instrument_type=instrument_type, instrument_id=instrument_id,
                    timestamp=observed_at.isoformat(), metric=metric, value=value, unit=unit,
                    source_file=relative, source_sheet=worksheet.title,
                )
    finally:
        workbook.close()


def iter_dataset_readings(root: str | Path) -> Iterator[NormalizedReading]:
    root = Path(root)
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}:
            yield from iter_workbook_readings(path, root=root)
            yield from iter_wide_readings(path, root=root)


def iter_wide_readings(path: str | Path, root: str | Path | None = None) -> Iterator[NormalizedReading]:
    """Parse reviewed combined BM and inclinometer summary tables."""
    path = Path(path)
    root_path = Path(root) if root else None
    relative = str(path.relative_to(root_path)) if root_path and path.is_relative_to(root_path) else path.name
    upper_name = path.name.upper()
    if "IN01" not in upper_name and "IN02" not in upper_name and " - BM " not in upper_name and "PIT - BM" not in upper_name:
        return
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "IN01" in upper_name or "IN02" in upper_name:
            worksheet = workbook["特征值表"]
            rows = list(worksheet.iter_rows(max_col=9, values_only=True))
            structure = "POWER_HOUSE" if "PH" in upper_name else "POWER_INTAKE"
            instrument_id = re.match(r"IN\d+(?:PH|PI)", upper_name).group(0)
            for row in rows[7:]:
                observed_at = _timestamp(row[0])
                value = _numeric(row[7] if len(row) > 7 else None) or _numeric(row[6] if len(row) > 6 else None)
                if observed_at and value is not None:
                    yield NormalizedReading(
                        structure=structure, instrument_type="IN", instrument_id=instrument_id,
                        timestamp=observed_at.isoformat(), metric="Resultant cumulative displacement", value=value,
                        unit="mm", source_file=relative, source_sheet=worksheet.title,
                    )
            return
        structure = "ACCRD" if "ACCRD" in upper_name else "POWER_HOUSE" if "GPH" in upper_name else "POWER_INTAKE"
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(max_row=worksheet.max_row, max_col=min(worksheet.max_column, 120), values_only=True))
            if len(rows) < 6:
                continue
            instrument_row = rows[1]
            settlement_start = next(
                (column for column, value in enumerate(rows[0]) if "沉降量" in _text(value) or "settlement" in _text(value).lower()),
                None,
            )
            if settlement_start is None:
                continue
            for column in range(settlement_start, len(instrument_row)):
                instrument_id = _text(instrument_row[column])
                if not re.match(r"^BM\d+(?:AD|GPH|PIT)$", instrument_id, re.IGNORECASE):
                    continue
                for row in rows[2:]:
                    observed_at = _timestamp(row[0] if row else None)
                    value = _numeric(row[column] if column < len(row) else None)
                    if observed_at and value is not None:
                        yield NormalizedReading(
                            structure=structure, instrument_type="BM", instrument_id=instrument_id,
                            timestamp=observed_at.isoformat(), metric="Cumulative settlement", value=value,
                            unit="mm", source_file=relative, source_sheet=worksheet.title,
                        )
    finally:
        workbook.close()


def write_jsonl(readings: Iterable[NormalizedReading], output: str | Path) -> dict[str, int]:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    counts: dict[str, int] = {}
    with output.open("w", encoding="utf-8") as handle:
        for reading in readings:
            handle.write(json.dumps(asdict(reading), ensure_ascii=False) + "\n")
            counts[reading.structure] = counts.get(reading.structure, 0) + 1
    return counts
